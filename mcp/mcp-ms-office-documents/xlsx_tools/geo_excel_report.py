"""Downloadable Excel export for the adaptive geo tool (render_geo_analysis).

Builds a 3-sheet workbook:
  * "Map"        — a picture of the exact map shown in chat (a server-side
                    screenshot). Static — editing data does not redraw it.
  * "Live Chart" — a NATIVE Excel Bubble Chart (dots positioned by real
                    longitude/latitude, sized by value) wired to real cell
                    ranges on the same sheet. Genuinely live: edit a value or
                    coordinate in those cells and the chart updates in Excel.
                    Only locations that resolved to a real coordinate appear
                    here — never invented, consistent with the chat tool.
  * "Data"       — the original uploaded rows, fully editable, untouched.

Bubble charts have no native per-point category color in openpyxl without
much more code, so this ships a single-color series (sized by value); a
"Top <group>" column is still included in the chart's source table for
reference when the analysis used a categorical/group breakdown.
"""
import io
import os
import uuid
import logging
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BubbleChart, Reference, Series
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

OUTPUT_DIR = "/app/output"

_TITLE_FONT = Font(size=16, bold=True, color="2F5597")
_SUB_FONT = Font(size=11, color="666666")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")


_CELL_SAFE_TYPES = (type(None), str, int, float, bool)


def _cell_safe(value):
    """openpyxl raises ValueError on cell values it can't natively store (a
    nested list/dict, for instance). Feature/row data arrives as arbitrary
    client-supplied JSON, so coerce anything unexpected to a string rather
    than let one odd value crash the whole export."""
    if isinstance(value, _CELL_SAFE_TYPES):
        return value
    return str(value)


def _autosize(ws, max_col: int, min_width: int = 10, max_width: int = 40) -> None:
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        longest = max(
            (len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row + 1)),
            default=min_width,
        )
        ws.column_dimensions[letter].width = max(min_width, min(max_width, longest + 2))


def _write_header_row(ws, row: int, headers: List[str]) -> None:
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center")


def _add_map_sheet(wb, title: str, why: str, currency_note: str,
                   warnings: List[str], map_image_bytes: Optional[bytes]) -> None:
    ws = wb.active
    ws.title = "Map"
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    sub = why + (f"  •  {currency_note}" if currency_note else "")
    ws.cell(row=2, column=1, value=sub).font = _SUB_FONT
    row = 4
    if map_image_bytes:
        try:
            img = XLImage(io.BytesIO(map_image_bytes))
            ws.add_image(img, f"A{row}")
        except Exception:
            logger.exception("Failed to embed map picture; continuing without it")
            ws.cell(row=row, column=1,
                    value="(Map picture unavailable — see the Live Chart sheet.)").font = _SUB_FONT
    else:
        ws.cell(row=row, column=1,
                value="(Map picture unavailable — see the Live Chart sheet.)").font = _SUB_FONT
    if warnings:
        warn_row = row + 34  # clear of the embedded picture in the common case
        ws.cell(row=warn_row, column=1, value="Warnings").font = Font(bold=True, color="8A6D00")
        for i, w in enumerate(warnings, start=1):
            ws.cell(row=warn_row + i, column=1, value=f"• {w}")


def _add_live_chart_sheet(wb, features: List[dict], location_field: str,
                          value_field: str, group_field: Optional[str]) -> None:
    """Native, editable Bubble Chart: x=longitude, y=latitude, size=value.
    Only features with a resolved centroid are included — a location that
    couldn't be matched to a real coordinate is never plotted with a made-up
    position (matches the chat tool's "never invent data" rule)."""
    def _is_num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool)

    ws = wb.create_sheet(title="Live Chart")
    plottable = [f for f in features
                if isinstance(f.get("centroid"), dict)
                and _is_num(f["centroid"].get("lat"))
                and _is_num(f["centroid"].get("lon"))]
    skipped = len(features) - len(plottable)

    headers = [location_field or "Location", "Longitude", "Latitude", value_field or "Value"]
    if group_field:
        headers.append(f"Top {group_field}")
    header_row = 1
    _write_header_row(ws, header_row, headers)

    for i, f in enumerate(plottable, start=1):
        r = header_row + i
        c = f["centroid"]
        ws.cell(row=r, column=1, value=_cell_safe(f.get("label") or f.get("location")))
        ws.cell(row=r, column=2, value=_cell_safe(c["lon"]))
        ws.cell(row=r, column=3, value=_cell_safe(c["lat"]))
        ws.cell(row=r, column=4, value=_cell_safe(f.get("value")))
        if group_field:
            ws.cell(row=r, column=5, value=_cell_safe(f.get("winner")))

    note_row = header_row + len(plottable) + 2
    note = ("Edit the Longitude / Latitude / " + (value_field or "Value") +
            " cells above and the chart updates automatically.")
    if skipped:
        note += (f" {skipped} location(s) had no resolvable coordinate and are not "
                 f"plotted here (see the Data sheet for everything).")
    ws.cell(row=note_row, column=1, value=note).font = _SUB_FONT

    if not plottable:
        _autosize(ws, len(headers))
        return

    n = len(plottable)
    chart = BubbleChart()
    chart.style = 18
    chart.title = "Locations (bubble size = " + (value_field or "value") + ")"
    chart.height = 10
    chart.width = 20
    xvalues = Reference(ws, min_col=2, min_row=header_row + 1, max_row=header_row + n)
    yvalues = Reference(ws, min_col=3, min_row=header_row + 1, max_row=header_row + n)
    size = Reference(ws, min_col=4, min_row=header_row + 1, max_row=header_row + n)
    series = Series(yvalues, xvalues, zvalues=size, title=value_field or "Value")
    chart.series.append(series)
    ws.add_chart(chart, f"{get_column_letter(len(headers) + 2)}{header_row}")
    _autosize(ws, len(headers))


def _add_data_sheet(wb, raw_rows: List[dict]) -> None:
    ws = wb.create_sheet(title="Data")
    if not raw_rows:
        ws.cell(row=1, column=1, value="(no source rows available)")
        return
    # Stable, readable column order: keys as first encountered across rows.
    seen, cols = set(), []
    for row in raw_rows:
        if not isinstance(row, dict):
            continue
        for k in row.keys():
            if k not in seen:
                seen.add(k)
                cols.append(k)
    _write_header_row(ws, 1, [str(c) for c in cols])
    r = 2
    for row in raw_rows:
        if not isinstance(row, dict):
            continue
        for j, col in enumerate(cols, start=1):
            ws.cell(row=r, column=j, value=_cell_safe(row.get(col)))
        r += 1
    _autosize(ws, len(cols))


def create_geo_excel_report(
    features: List[dict],
    raw_rows: List[dict],
    location_field: str,
    value_field: str,
    group_field: Optional[str],
    title: str,
    why: str,
    currency_note: str,
    warnings: List[str],
    map_image_bytes: Optional[bytes] = None,
) -> str:
    """Build the 3-sheet geo workbook and return its public download URL."""
    logger.info("Creating geo Excel report: %d features, %d raw rows, image=%s",
                len(features), len(raw_rows), bool(map_image_bytes))
    wb = Workbook()
    _add_map_sheet(wb, title, why, currency_note, warnings, map_image_bytes)
    _add_live_chart_sheet(wb, features, location_field, value_field, group_field)
    _add_data_sheet(wb, raw_rows)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"geo_report_{uuid.uuid4().hex}.xlsx"
    file_path = os.path.join(OUTPUT_DIR, filename)
    wb.save(file_path)
    public_url = f"{os.getenv('PUBLIC_FILES_BASE_URL', 'http://10.10.30.160:7001/files').rstrip('/')}/{filename}"
    logger.info("Geo Excel report saved: %s", public_url)
    return public_url
