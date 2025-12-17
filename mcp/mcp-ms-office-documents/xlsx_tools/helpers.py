import re
import logging
from typing import List, Tuple, Dict, Optional
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def parse_table(lines: List[str], start_idx: int) -> Tuple[Optional[List[List[str]]], int]:
    """Parse markdown table and return (table_data, next_index)."""
    table_lines: List[str] = []
    i = start_idx

    # Find all consecutive table lines
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
        else:
            break

    if len(table_lines) < 2:  # Need at least header and separator
        return None, start_idx + 1

    # Parse table data skipping separator row
    table_data: List[List[str]] = []
    for line in table_lines:
        if '---' in line or ':-:' in line or ':--' in line or '--:' in line:
            continue
        cells = [cell.strip() for cell in line.split('|')[1:-1]]
        table_data.append(cells)

    return table_data, i


def format_cell_value(value: str):
    """Convert string value to appropriate Excel type (number, text, formula, etc.)."""
    value = value.strip()

    # Check if it's a formula (starts with =)
    if value.startswith('='):
        return value

    # Try to convert to number
    try:
        if value.endswith('%'):
            return float(value[:-1]) / 100
        return float(value)
    except ValueError:
        return value


def parse_cell_formatting(cell_text: str) -> Tuple[str, Dict[str, bool]]:
    """Parse markdown formatting in cell text and return clean text and formatting info."""
    formatting_info = {'bold': False, 'italic': False, 'monospace': False}
    clean_text = cell_text.strip()

    # Bold **text**
    if clean_text.startswith('**') and clean_text.endswith('**'):
        clean_text = clean_text[2:-2]
        formatting_info['bold'] = True
    # Italic *text*
    elif clean_text.startswith('*') and clean_text.endswith('*'):
        clean_text = clean_text[1:-1]
        formatting_info['italic'] = True
    # Monospace `text`
    elif clean_text.startswith('`') and clean_text.endswith('`'):
        clean_text = clean_text[1:-1]
        formatting_info['monospace'] = True

    return clean_text, formatting_info


def apply_cell_formatting(cell, formatting_info: Dict[str, bool]) -> None:
    """Apply formatting information to an Excel cell."""
    current_font = cell.font
    if formatting_info['bold']:
        cell.font = Font(bold=True, color=current_font.color, size=current_font.size)
    elif formatting_info['italic']:
        cell.font = Font(italic=True, color=current_font.color, size=current_font.size)
    elif formatting_info['monospace']:
        cell.font = Font(name='Courier New', color=current_font.color, size=current_font.size)


def adjust_formula_references(formula: str, current_excel_row: int, table_positions: Optional[Dict[str, int]] = None) -> str:
    """Convert row-relative references [offset] and table references T1.B[1] to actual Excel row numbers."""
    if not formula.startswith('='):
        return formula

    if table_positions is None:
        table_positions = {}

    # Table cell references e.g. T1.B[1]
    table_pattern = r'T(\d+)\.([A-Z]+)\[([+-]?\d+)\]'

    def replace_table_reference(match):
        table_num = int(match.group(1))
        column = match.group(2)
        offset = int(match.group(3))
        table_key = f"T{table_num}"
        if table_key in table_positions:
            table_start_row = table_positions[table_key]
            actual_row = table_start_row + 1 + offset
            return f"{column}{actual_row}"
        actual_row = current_excel_row + offset
        return f"{column}{actual_row}"

    adjusted = re.sub(table_pattern, replace_table_reference, formula)

    # Table range references e.g. T1.B[0]:T1.E[0]
    table_range_pattern = r'T(\d+)\.([A-Z]+)\([+-]?\d+\):T(\d+)\.([A-Z]+)\([+-]?\d+\)'

    # Corrected with brackets
    table_range_pattern = r'T(\d+)\.([A-Z]+)\[([+-]?\d+)\]:T(\d+)\.([A-Z]+)\[([+-]?\d+)\]'

    def replace_table_range(match):
        start_table_num = int(match.group(1))
        start_col = match.group(2)
        start_offset = int(match.group(3))
        end_table_num = int(match.group(4))
        end_col = match.group(5)
        end_offset = int(match.group(6))

        start_key = f"T{start_table_num}"
        end_key = f"T{end_table_num}"

        if start_key in table_positions:
            start_table_row = table_positions[start_key]
            start_row = start_table_row + 1 + start_offset
        else:
            start_row = current_excel_row + start_offset

        if end_key in table_positions:
            end_table_row = table_positions[end_key]
            end_row = end_table_row + 1 + end_offset
        else:
            end_row = current_excel_row + end_offset

        return f"{start_col}{start_row}:{end_col}{end_row}"

    adjusted = re.sub(table_range_pattern, replace_table_range, adjusted)

    # Simplified function over table range e.g. T1.SUM(B[0]:E[0])
    table_func_pattern = r'T(\d+)\.(SUM|AVERAGE|MAX|MIN)\(([A-Z]+)\[([+-]?\d+)\]:([A-Z]+)\[([+-]?\d+)\]\)'

    def replace_table_function(match):
        table_num = int(match.group(1))
        func_name = match.group(2)
        start_col = match.group(3)
        start_offset = int(match.group(4))
        end_col = match.group(5)
        end_offset = int(match.group(6))

        key = f"T{table_num}"
        if key in table_positions:
            table_start_row = table_positions[key]
            start_row = table_start_row + 1 + start_offset
            end_row = table_start_row + 1 + end_offset
        else:
            start_row = current_excel_row + start_offset
            end_row = current_excel_row + end_offset

        return f"={func_name}({start_col}{start_row}:{end_col}{end_row})"

    adjusted = re.sub(table_func_pattern, replace_table_function, adjusted)

    # Determine current table start for relative references
    current_table_start = None
    for table_key, table_start_row in table_positions.items():
        if table_start_row <= current_excel_row:
            current_table_start = table_start_row

    # Handle row-relative references e.g. B[0]
    rel_pattern = r'([A-Z]+)\[([+-]?\d+)\]'

    def replace_rel(match):
        column = match.group(1)
        offset = int(match.group(2))
        if current_table_start is not None:
            actual_row = current_table_start + 1 + offset
        else:
            actual_row = current_excel_row + offset
        return f"{column}{actual_row}"

    adjusted = re.sub(rel_pattern, replace_rel, adjusted)

    # Row-relative range e.g. B[0]:E[0]
    range_pattern = r'([A-Z]+)\[([+-]?\d+)\]:([A-Z]+)\[([+-]?\d+)\]'

    def replace_range(match):
        start_col = match.group(1)
        start_offset = int(match.group(2))
        end_col = match.group(3)
        end_offset = int(match.group(4))
        if current_table_start is not None:
            start_row = current_table_start + 1 + start_offset
            end_row = current_table_start + 1 + end_offset
        else:
            start_row = current_excel_row + start_offset
            end_row = current_excel_row + end_offset
        return f"{start_col}{start_row}:{end_col}{end_row}"

    adjusted = re.sub(range_pattern, replace_range, adjusted)

    return adjusted


def detect_formula_pattern(value: str) -> str:
    """Detect common formula patterns in markdown and convert to Excel formulas."""
    value = value.strip()
    if value.startswith('='):
        return value
    if re.match(r'^(SUM|sum)\([A-Z]+\d+:[A-Z]+\d+\)$', value):
        return f"={value.upper()}"
    if re.match(r'^(AVG|avg|AVERAGE|average)\([A-Z]+\d+:[A-Z]+\d+\)$', value):
        return f"=AVERAGE({value.split('(')[1]}"
    if re.match(r'^[A-Z]+\d+[+\-*/][A-Z]+\d+$', value):
        return f"={value}"
    if re.match(r'^[A-Z]+\d+/[A-Z]+\d+\*100$', value):
        return f"={value}/100"
    return value


def add_table_to_sheet(
    table_data: List[List[str]],
    worksheet,
    start_row: int,
    table_positions: Optional[Dict[str, int]] = None
) -> int:
    """Add table data to Excel worksheet with proper formatting and formula support."""
    if not table_data:
        return start_row

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    formula_fill = PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Fill cells
    for row_idx, row_data in enumerate(table_data):
        current_excel_row = start_row + row_idx
        for col_idx, cell_text in enumerate(row_data):
            cell = worksheet.cell(row=current_excel_row, column=col_idx + 1)
            clean_text, formatting_info = parse_cell_formatting(cell_text)
            formula_value = detect_formula_pattern(clean_text)

            if isinstance(formula_value, str) and formula_value.startswith('='):
                adjusted_formula = adjust_formula_references(formula_value, current_excel_row, table_positions)
                cell.value = adjusted_formula
                cell.fill = formula_fill
            else:
                formatted_value = format_cell_value(clean_text)
                cell.value = formatted_value

            apply_cell_formatting(cell, formatting_info)
            cell.border = border

            # Alignment and number formats
            if row_idx == 0:
                cell.alignment = Alignment(horizontal='center')
            elif isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith('=')):
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='left')

            if row_idx == 0:
                cell.font = header_font
                cell.fill = header_fill
            elif isinstance(cell.value, float) and 0 < cell.value <= 1:
                cell.number_format = '0.00%'
            elif isinstance(cell.value, (int, float)) and cell.value >= 1000:
                cell.number_format = '#,##0'

    # Column widths
    for col_idx in range(len(table_data[0]) if table_data else 0):
        column_letter = get_column_letter(col_idx + 1)
        max_length = 0
        for row in table_data:
            if col_idx < len(row):
                max_length = max(max_length, len(str(row[col_idx])))
        adjusted_width = min(max(max_length + 2, 12), 25)
        worksheet.column_dimensions[column_letter].width = adjusted_width

    return start_row + len(table_data) + 2
