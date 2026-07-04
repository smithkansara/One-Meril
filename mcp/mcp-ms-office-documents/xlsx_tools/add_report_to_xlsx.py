import io
import os
import re
import uuid
import logging
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, PieChart, DoughnutChart, RadarChart, Reference

from .helpers import parse_table, add_table_to_sheet

logger = logging.getLogger(__name__)

OUTPUT_DIR = "/app/output"


def _populate_sheet_from_markdown(ws, markdown_content: str) -> None:
    """Parse markdown (headers, tables, plain text) and populate a worksheet."""
    lines: List[str] = markdown_content.split('\n')
    current_row = 1
    table_counter = 1
    table_positions = {}
    i = 0

    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        if line.startswith('#'):
            header_level = len(line) - len(line.lstrip('#'))
            header_text = line.lstrip('#').strip()
            cell = ws.cell(row=current_row, column=1)
            cell.value = header_text
            if header_level == 1:
                cell.font = Font(size=16, bold=True, color="2F5597")
            elif header_level == 2:
                cell.font = Font(size=14, bold=True, color="4472C4")
            else:
                cell.font = Font(size=12, bold=True)
            current_row += 2
            i += 1

        elif line.startswith('|'):
            table_data, i = parse_table(lines, i)
            if table_data:
                table_key = f"T{table_counter}"
                table_positions[table_key] = current_row
                current_row = add_table_to_sheet(table_data, ws, current_row, table_positions)
                table_counter += 1
        else:
            cell = ws.cell(row=current_row, column=1)
            cell.value = line
            current_row += 1
            i += 1


def _coerce_number(value):
    """Best-effort convert a Chart.js data value to a float (strip commas/currency)."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        stripped = re.sub(r"[^0-9.\-]", "", value)
        try:
            return float(stripped)
        except ValueError:
            return None
    return None


def _make_chart_object(chart_type: str, index_axis: Optional[str]):
    """Map a Chart.js chart type to an openpyxl chart object."""
    ct = (chart_type or "bar").lower()
    if ct == "pie":
        return PieChart()
    if ct in ("doughnut", "donut"):
        return DoughnutChart()
    if ct == "line":
        return LineChart()
    if ct == "radar":
        chart = RadarChart()
        chart.style = 26
        return chart
    # default: bar (vertical columns, or horizontal bars when indexAxis == "y")
    chart = BarChart()
    chart.type = "bar" if index_axis == "y" else "col"
    return chart


def _add_charts_sheet(wb, charts: List[dict], sheet_name: str = "Charts",
                      heading: str = "") -> None:
    """Render each Chart.js config as a native Excel chart on a dedicated sheet.

    Writes each chart's data as a small table (category column + one column per
    dataset), then anchors a native openpyxl chart beside it. Charts are stacked
    vertically so multiple charts don't overlap. An optional heading naming the
    requested analysis is shown at the top; the charts follow below it.
    """
    ws = wb.create_sheet(title=sheet_name[:31], index=1)
    ws.column_dimensions["A"].width = 28
    anchor_row = 1

    # Research/analysis heading at the very top; charts start below it.
    if heading:
        cell = ws.cell(row=anchor_row, column=1, value=heading)
        cell.font = Font(size=16, bold=True, color="2F5597")
        anchor_row += 2

    for chart in charts:
        config = chart.get("config") or {}
        data = config.get("data") or {}
        labels = data.get("labels") or []
        datasets = data.get("datasets") or []
        if not labels or not datasets:
            continue

        options = config.get("options") or {}
        title = (
            chart.get("title")
            or options.get("plugins", {}).get("title", {}).get("text")
            or "Chart"
        )
        index_axis = options.get("indexAxis")
        n = len(labels)

        # --- title + data table ---
        ws.cell(row=anchor_row, column=1, value=title).font = Font(size=12, bold=True, color="2F5597")
        header_row = anchor_row + 1
        ws.cell(row=header_row, column=1, value="Category").font = Font(bold=True)
        for j, ds in enumerate(datasets):
            ws.cell(row=header_row, column=2 + j, value=ds.get("label") or f"Series {j + 1}").font = Font(bold=True)
        for r, label in enumerate(labels):
            ws.cell(row=header_row + 1 + r, column=1, value=str(label))
            for j, ds in enumerate(datasets):
                vals = ds.get("data") or []
                ws.cell(row=header_row + 1 + r, column=2 + j,
                        value=_coerce_number(vals[r]) if r < len(vals) else None)

        # --- native chart ---
        chart_obj = _make_chart_object(config.get("type", "bar"), index_axis)
        chart_obj.title = title
        max_col = 1 + len(datasets)
        data_ref = Reference(ws, min_col=2, max_col=max_col, min_row=header_row, max_row=header_row + n)
        cats_ref = Reference(ws, min_col=1, max_col=1, min_row=header_row + 1, max_row=header_row + n)
        chart_obj.add_data(data_ref, titles_from_data=True)
        chart_obj.set_categories(cats_ref)
        chart_obj.height = 8   # cm
        chart_obj.width = 16   # cm
        ws.add_chart(chart_obj, f"{get_column_letter(max_col + 2)}{anchor_row}")

        # leave room for the chart image (~16 rows) or the data table, whichever is taller
        anchor_row = header_row + max(n + 3, 18)


def create_excel_with_report(
    data_content: str,
    report_markdown: str,
    data_sheet_name: str = "Data",
    report_sheet_name: str = "Analysis Report",
    charts: Optional[List[dict]] = None,
    charts_heading: str = "",
) -> str:
    """Create a multi-sheet Excel: 'Data' sheet with original content, an optional
    'Charts' sheet with native charts, and an 'Analysis Report' sheet with the report.

    Returns the public download URL.
    """
    logger.info("Creating Excel workbook with '%s' + '%s' sheets (%d charts)",
                data_sheet_name, report_sheet_name, len(charts or []))

    wb = Workbook()

    # First sheet: original data
    ws_data = wb.active
    ws_data.title = data_sheet_name[:31]
    _populate_sheet_from_markdown(ws_data, data_content)

    # Second sheet (optional): native charts under a research heading
    if charts:
        try:
            _add_charts_sheet(wb, charts, heading=charts_heading)
        except Exception:
            logger.exception("Failed to render charts sheet; continuing without charts")

    # Last sheet: analysis report
    ws_report = wb.create_sheet(title=report_sheet_name[:31])
    _populate_sheet_from_markdown(ws_report, report_markdown)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"analysis_{uuid.uuid4().hex}.xlsx"
    file_path = os.path.join(OUTPUT_DIR, filename)

    try:
        wb.save(file_path)
        public_url = f"{os.getenv('PUBLIC_FILES_BASE_URL', 'http://10.10.30.160:7001/files').rstrip('/')}/{filename}"
        logger.info("Excel with report saved: %s", public_url)
        return public_url
    except Exception as e:
        logger.error("Error saving Excel with report: %s", str(e), exc_info=True)
        return f"Error saving Excel with report: {str(e)}"
