"""
Unified report builder: combines original source data + analysis report
into a single downloadable file in the same format as the source.

Supports: xlsx, docx, pdf, pptx, csv, txt
"""

import io
import logging
import os
import re
import uuid
from typing import List, Tuple, Dict

logger = logging.getLogger(__name__)

OUTPUT_DIR = "/app/output"
PUBLIC_BASE = "https://chatgpt.e-meril.in/files"


def _save(buffer_or_path, suffix: str) -> str:
    """Save bytes/BytesIO to OUTPUT_DIR and return public URL."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"report_{uuid.uuid4().hex}.{suffix}"
    file_path = os.path.join(OUTPUT_DIR, filename)
    if isinstance(buffer_or_path, (bytes, bytearray)):
        with open(file_path, "wb") as f:
            f.write(buffer_or_path)
    elif hasattr(buffer_or_path, "read"):
        buffer_or_path.seek(0)
        with open(file_path, "wb") as f:
            f.write(buffer_or_path.read())
    else:
        # already a path — just rename/copy isn't needed, caller saved to this path
        file_path = buffer_or_path
        filename = os.path.basename(file_path)
    return f"{PUBLIC_BASE}/{filename}"


# ---------------------------------------------------------------------------
# XLSX  — two sheets: Data  +  Report
# ---------------------------------------------------------------------------

def _build_xlsx(data_content: str, report_markdown: str,
                data_sheet: str, report_sheet: str) -> str:
    from xlsx_tools.add_report_to_xlsx import create_excel_with_report
    return create_excel_with_report(
        data_content, report_markdown,
        data_sheet_name=data_sheet,
        report_sheet_name=report_sheet,
    )


# ---------------------------------------------------------------------------
# DOCX  — data section + page-break + report section
# ---------------------------------------------------------------------------

def _build_docx(data_content: str, report_markdown: str,
                data_title: str, report_title: str) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx_tools.helpers import (
        load_templates, parse_inline_formatting,
        parse_table, add_table_to_doc, process_list_items,
    )

    path = load_templates()
    doc = Document(path) if path else Document()

    def _render(markdown_content: str):
        lines = markdown_content.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i]
            if not line.strip():
                i += 1
                continue
            line_s = line.strip()
            if line_s.startswith('#'):
                lvl = len(line_s) - len(line_s.lstrip('#'))
                heading = doc.add_heading('', level=min(lvl, 6))
                parse_inline_formatting(line_s.lstrip('#').strip(), heading)
                i += 1
            elif line_s.startswith('|'):
                table_data, i = parse_table(lines, i)
                if table_data:
                    add_table_to_doc(table_data, doc)
            elif re.match(r'^\d+\.\s+', line_s):
                i = process_list_items(lines, i, doc, True, 0)
            elif re.match(r'^[-*+]\s+', line_s):
                i = process_list_items(lines, i, doc, False, 0)
            elif line_s.startswith('>'):
                p = doc.add_paragraph()
                p.style = 'Quote'
                parse_inline_formatting(line_s[1:].strip(), p)
                i += 1
            else:
                p = doc.add_paragraph()
                parse_inline_formatting(line_s, p)
                i += 1

    # --- Data section ---
    h = doc.add_heading(data_title, level=1)
    _render(data_content)

    # Page break
    doc.add_page_break()

    # --- Report section ---
    doc.add_heading(report_title, level=1)
    _render(report_markdown)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"report_{uuid.uuid4().hex}.docx"
    file_path = os.path.join(OUTPUT_DIR, filename)
    doc.save(file_path)
    logger.info("DOCX report saved: %s", file_path)
    return f"{PUBLIC_BASE}/{filename}"


# ---------------------------------------------------------------------------
# PDF  — data section + page-break + report section
# ---------------------------------------------------------------------------

def _build_pdf(data_content: str, report_markdown: str,
               data_title: str, report_title: str,
               page_size: str = "letter") -> str:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_JUSTIFY
    from pdf_tools.helpers import parse_inline_formatting_pdf, parse_table as pdf_parse_table, process_list_items_pdf

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"report_{uuid.uuid4().hex}.pdf"
    file_path = os.path.join(OUTPUT_DIR, filename)

    pagesize = letter if page_size == "letter" else A4
    doc_obj = SimpleDocTemplate(
        file_path, pagesize=pagesize,
        rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('CH1', parent=styles['Heading1'],
                              fontSize=18, textColor=colors.HexColor('#2c2c2c'),
                              spaceAfter=12, spaceBefore=12))
    styles.add(ParagraphStyle('CH2', parent=styles['Heading2'],
                              fontSize=16, textColor=colors.HexColor('#2c2c2c'),
                              spaceAfter=10, spaceBefore=10))
    styles.add(ParagraphStyle('CH3', parent=styles['Heading3'],
                              fontSize=14, textColor=colors.HexColor('#2c2c2c'),
                              spaceAfter=8, spaceBefore=8))
    styles.add(ParagraphStyle('CB', parent=styles['BodyText'],
                              fontSize=11, leading=14,
                              alignment=TA_JUSTIFY, spaceAfter=12))
    styles.add(ParagraphStyle('CQ', parent=styles['BodyText'],
                              fontSize=11, leading=14,
                              leftIndent=36, rightIndent=36,
                              textColor=colors.HexColor('#555555'), spaceAfter=12))

    story = []

    def _render(markdown_content: str):
        lines = markdown_content.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i]
            if not line.strip():
                story.append(Spacer(1, 0.15 * inch))
                i += 1
                continue
            line_s = line.strip()
            if line_s.startswith('#'):
                lvl = len(line_s) - len(line_s.lstrip('#'))
                text = parse_inline_formatting_pdf(line_s.lstrip('#').strip())
                sty = styles.get('CH1' if lvl == 1 else 'CH2' if lvl == 2 else 'CH3')
                story.append(Paragraph(text, sty))
                i += 1
            elif line_s.startswith('|'):
                tdata, i = pdf_parse_table(lines, i)
                if tdata:
                    t = Table(tdata)
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ]))
                    story.append(t)
                    story.append(Spacer(1, 0.2 * inch))
            elif re.match(r'^\d+\.\s+', line_s) or re.match(r'^[-*+]\s+', line_s):
                is_ordered = bool(re.match(r'^\d+\.\s+', line_s))
                items, i = process_list_items_pdf(lines, i, is_ordered)
                for item_text, level in items:
                    fmt = parse_inline_formatting_pdf(item_text)
                    li_style = ParagraphStyle('LI', parent=styles['CB'],
                                             leftIndent=level * 20,
                                             bulletIndent=level * 20 - 10)
                    story.append(Paragraph(f"• {fmt}", li_style))
                story.append(Spacer(1, 0.1 * inch))
            elif line_s.startswith('>'):
                fmt = parse_inline_formatting_pdf(line_s[1:].strip())
                story.append(Paragraph(fmt, styles['CQ']))
                i += 1
            else:
                fmt = parse_inline_formatting_pdf(line_s)
                story.append(Paragraph(fmt, styles['CB']))
                i += 1

    # Section title helper
    def section_heading(title: str):
        story.append(Paragraph(
            f'<font size="20"><b>{title}</b></font>',
            ParagraphStyle('SH', parent=styles['CB'],
                           textColor=colors.HexColor('#1a3c6e'),
                           spaceAfter=18, spaceBefore=6)
        ))

    section_heading(data_title)
    _render(data_content)
    story.append(PageBreak())
    section_heading(report_title)
    _render(report_markdown)

    doc_obj.build(story)
    logger.info("PDF report saved: %s", file_path)
    return f"{PUBLIC_BASE}/{filename}"


# ---------------------------------------------------------------------------
# PPTX  — data section slides + report section slides
# ---------------------------------------------------------------------------

def _markdown_to_slides(markdown_content: str, section_title: str) -> List[Dict]:
    """Convert markdown to a list of slide dicts usable by PowerpointPresentation."""
    slides = []
    # Section divider slide
    slides.append({"slide_type": "section", "slide_title": section_title})

    lines = markdown_content.split('\n')
    current_slide_title = ""
    current_bullets: List[Dict] = []

    def flush():
        if current_slide_title or current_bullets:
            slides.append({
                "slide_type": "content",
                "slide_title": current_slide_title or "Details",
                "slide_text": list(current_bullets) if current_bullets else None,
            })

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if line.startswith('## ') or line.startswith('### '):
            flush()
            current_slide_title = line.lstrip('#').strip()
            current_bullets = []
            i += 1
        elif line.startswith('# '):
            # Top-level heading → section slide
            flush()
            slides.append({"slide_type": "section", "slide_title": line.lstrip('#').strip()})
            current_slide_title = ""
            current_bullets = []
            i += 1
        elif line.startswith('|'):
            # Table → skip (tables don't render well in PPTX bullets)
            while i < len(lines) and lines[i].strip().startswith('|'):
                i += 1
        elif re.match(r'^[-*+]\s+', line) or re.match(r'^\d+\.\s+', line):
            text = re.sub(r'^[-*+]\s+', '', line)
            text = re.sub(r'^\d+\.\s+', '', text)
            current_bullets.append({"text": text, "indentation_level": 1})
            i += 1
        else:
            if line and not line.startswith('---'):
                current_bullets.append({"text": line, "indentation_level": 1})
            i += 1

    flush()
    return slides


def _build_pptx(data_content: str, report_markdown: str,
                data_title: str, report_title: str,
                fmt: str = "4:3") -> str:
    from pptx_tools.helpers import PowerpointPresentation

    data_slides = _markdown_to_slides(data_content, data_title)
    report_slides = _markdown_to_slides(report_markdown, report_title)
    all_slides = data_slides + report_slides

    if not all_slides:
        all_slides = [{"slide_type": "title", "slide_title": "Report", "author": ""}]

    presentation = PowerpointPresentation(all_slides, fmt)
    file_path = presentation.save()
    filename = os.path.basename(file_path)
    logger.info("PPTX report saved: %s", file_path)
    return f"{PUBLIC_BASE}/{filename}"


# ---------------------------------------------------------------------------
# CSV  — original rows + separator + report as table rows / plain text
# ---------------------------------------------------------------------------

def _build_csv(data_content: str, report_markdown: str,
               data_title: str, report_title: str) -> str:
    lines = []
    lines.append(f"# {data_title}")
    lines.append(data_content.strip())
    lines.append("")
    lines.append("")
    lines.append(f"# {report_title}")
    # Convert markdown tables to CSV-style; keep plain lines as-is
    for line in report_markdown.split('\n'):
        s = line.strip()
        if s.startswith('|') and s.endswith('|'):
            # strip separator rows
            if re.match(r'^[\|:\- ]+$', s):
                continue
            # strip leading/trailing pipes and join with comma
            cells = [c.strip() for c in s.strip('|').split('|')]
            lines.append(','.join(f'"{c}"' for c in cells))
        else:
            lines.append(line)

    content = '\n'.join(lines)
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"report_{uuid.uuid4().hex}.csv"
    file_path = os.path.join(OUTPUT_DIR, filename)
    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(content)
    logger.info("CSV report saved: %s", file_path)
    return f"{PUBLIC_BASE}/{filename}"


# ---------------------------------------------------------------------------
# TXT  — plain text with separator
# ---------------------------------------------------------------------------

def _build_txt(data_content: str, report_markdown: str,
               data_title: str, report_title: str) -> str:
    sep = "=" * 60
    content = (
        f"{sep}\n{data_title.upper()}\n{sep}\n\n"
        f"{data_content.strip()}\n\n"
        f"{sep}\n{report_title.upper()}\n{sep}\n\n"
        f"{report_markdown.strip()}\n"
    )
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"report_{uuid.uuid4().hex}.txt"
    file_path = os.path.join(OUTPUT_DIR, filename)
    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(content)
    logger.info("TXT report saved: %s", file_path)
    return f"{PUBLIC_BASE}/{filename}"


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

_FORMAT_ALIASES = {
    "xls": "xlsx",
    "doc": "docx",
    "ppt": "pptx",
    "text": "txt",
    "plain": "txt",
}

SUPPORTED_FORMATS = {"xlsx", "docx", "pdf", "pptx", "csv", "txt"}


def create_report_in_same_format(
    data_content: str,
    report_markdown: str,
    file_format: str,
    data_section_title: str = "Source Data",
    report_section_title: str = "Analysis Report",
) -> str:
    """Combine original data + analysis report into a single file.

    Returns the public download URL.

    :param data_content:       Original file/web data as text or markdown.
    :param report_markdown:    Analysis report in markdown.
    :param file_format:        Target format: xlsx | docx | pdf | pptx | csv | txt.
    :param data_section_title: Label for the data section/sheet.
    :param report_section_title: Label for the report section/sheet.
    """
    fmt = file_format.lower().strip().lstrip('.')
    fmt = _FORMAT_ALIASES.get(fmt, fmt)

    if fmt not in SUPPORTED_FORMATS:
        logger.warning("Unknown format '%s', falling back to docx", fmt)
        fmt = "docx"

    logger.info(
        "Building combined report: format=%s data_len=%d report_len=%d",
        fmt, len(data_content), len(report_markdown),
    )

    if fmt == "xlsx":
        return _build_xlsx(data_content, report_markdown, data_section_title, report_section_title)
    elif fmt == "docx":
        return _build_docx(data_content, report_markdown, data_section_title, report_section_title)
    elif fmt == "pdf":
        return _build_pdf(data_content, report_markdown, data_section_title, report_section_title)
    elif fmt == "pptx":
        return _build_pptx(data_content, report_markdown, data_section_title, report_section_title)
    elif fmt == "csv":
        return _build_csv(data_content, report_markdown, data_section_title, report_section_title)
    else:  # txt
        return _build_txt(data_content, report_markdown, data_section_title, report_section_title)
