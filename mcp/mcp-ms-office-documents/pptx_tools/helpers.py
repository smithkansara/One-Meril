import io
import logging
from typing import List, Dict, Any, Tuple, Optional
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from template_utils import find_pptx_templates

# Slide layout constants
TITLE_LAYOUT = 2
SECTION_LAYOUT = 7
CONTENT_LAYOUT = 4

logger = logging.getLogger(__name__)


def load_templates() -> Tuple[Optional[str], Optional[str]]:
    """Resolve presentation templates (4:3, 16:9) from custom/default template dirs.

    Returns: tuple[str|None, str|None] -> (path_4_3, path_16_9)
    """
    t43, t169 = find_pptx_templates()
    if not t43 or not t169:
        logger.info("One or more PPT templates missing; will fall back to PowerPoint defaults where needed")
    return t43, t169


class PowerpointPresentation:
    """Helper to build a PPTX presentation from structured slide dictionaries."""

    def __init__(self, slides: List[Dict[str, Any]], format: str):
        """Initialize PowerPoint presentation with slides and format."""
        logger.info(f"Initializing PowerPoint: slides={len(slides)}, format={format}")

        # Validate input
        if not slides:
            raise ValueError("At least one slide is required")

        # Load templates
        self.template_regular, self.template_wide = load_templates()
        logger.debug(f"Selected templates -> 4:3={self.template_regular}, 16:9={self.template_wide}")

        # Create presentation based on the format used
        try:
            if format == "4:3":
                if self.template_regular:
                    self.presentation = Presentation(self.template_regular)
                else:
                    self.presentation = Presentation()  # Use default template
                    logger.warning("No 4:3 template found, using PowerPoint default template")
            elif format == "16:9":
                if self.template_wide:
                    self.presentation = Presentation(self.template_wide)
                else:
                    self.presentation = Presentation()  # Use default template
                    logger.warning("No 16:9 template found, using PowerPoint default template")
            else:
                logger.warning(f"Unknown format '{format}', defaulting to 4:3")
                if self.template_regular:
                    self.presentation = Presentation(self.template_regular)
                else:
                    self.presentation = Presentation()  # Use default template
        except Exception as e:
            logger.error(f"Failed to load template: {e}")
            logger.info("Falling back to default PowerPoint template")
            self.presentation = Presentation()  # Fallback to default template

        # Remove default slide if it exists (some templates add one automatically)
        if len(self.presentation.slides) > 0:
            try:
                logger.debug("Removing default first slide from new presentation")
                slide_to_remove = self.presentation.slides[0]
                # Use underlying element removal to clear the initial slide
                self.presentation.slides.element.remove(slide_to_remove.element)
            except Exception as e:
                logger.debug(f"Could not remove default slide (non-fatal): {e}")

        # Create slides
        self._create_slides(slides)

    def _create_slides(self, slides: List[Dict[str, Any]]):
        """Create all slides from the slides data."""
        logger.info(f"Creating {len(slides)} slides")
        for i, slide in enumerate(slides):
            try:
                slide_type = slide.get("slide_type")
                logger.debug(f"Creating slide index={i}, type={slide_type}, title={slide.get('slide_title','')}")

                if slide_type == "content":
                    self.create_content_slide(slide)
                elif slide_type == "section":
                    self.create_section_slide(slide)
                elif slide_type == "title":
                    self.create_title_slide(slide)
                else:
                    logger.warning(f"Unknown slide type '{slide_type}' for slide {i}, skipping")

            except Exception as e:
                logger.error(f"Failed to create slide {i}: {e}")
                raise ValueError(f"Error creating slide {i}: {str(e)}")

    def create_title_slide(self, slide: Dict[str, Any]):
        """Create a title slide."""
        try:
            title_layout = self.presentation.slide_layouts[TITLE_LAYOUT]
            title_slide = self.presentation.slides.add_slide(title_layout)
            logger.debug("Added title slide")

            # Set title
            if len(title_slide.placeholders) > 0:
                title_text = slide.get("slide_title", "")
                title_slide.placeholders[0].text = title_text
                logger.debug(f"Title slide title set: {title_text!r}")

            # Set author
            if len(title_slide.placeholders) > 1:
                author_text = slide.get("author", "")
                title_slide.placeholders[1].text = author_text
                logger.debug(f"Title slide author set: {author_text!r}")

        except Exception as e:
            logger.error(f"Failed to create title slide: {e}")
            raise

    def create_section_slide(self, slide: Dict[str, Any]):
        """Create a section slide."""
        try:
            section_layout = self.presentation.slide_layouts[SECTION_LAYOUT]
            section_slide = self.presentation.slides.add_slide(section_layout)
            logger.debug("Added section slide")

            # Set title
            if len(section_slide.placeholders) > 0:
                title_text = slide.get("slide_title", "")
                section_slide.placeholders[0].text = title_text
                logger.debug(f"Section slide title set: {title_text!r}")

        except Exception as e:
            logger.error(f"Failed to create section slide: {e}")
            raise

    def create_content_slide(self, slide: Dict[str, Any]):
        """Create a content slide with bullet points."""
        try:
            content_layout = self.presentation.slide_layouts[CONTENT_LAYOUT]
            content_slide = self.presentation.slides.add_slide(content_layout)
            logger.debug("Added content slide")

            # Set title
            if len(content_slide.placeholders) > 0:
                title_text = slide.get("slide_title", "")
                content_slide.placeholders[0].text = title_text
                logger.debug(f"Content slide title set: {title_text!r}")

            # Add content
            slide_text = slide.get("slide_text", [])
            if slide_text and len(content_slide.placeholders) > 1:
                logger.debug(f"Adding {len(slide_text)} bullet items to content slide")
                # Clear existing text
                content_slide.placeholders[1].text = ""

                # Add first paragraph
                first_item = slide_text[0]
                first_text = first_item.get("text", "")
                first_level = max(0, int(first_item.get("indentation_level", 1)) - 1)
                content_slide.placeholders[1].text_frame.paragraphs[0].text = first_text
                content_slide.placeholders[1].text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
                content_slide.placeholders[1].text_frame.paragraphs[0].level = first_level
                logger.debug(f"Bullet[0]: level={first_level} text={first_text!r}")

                # Add remaining paragraphs
                for idx, paragraph_data in enumerate(slide_text[1:], start=1):
                    p = content_slide.placeholders[1].text_frame.add_paragraph()
                    p.text = paragraph_data.get("text", "")
                    p.alignment = PP_ALIGN.LEFT
                    level = max(0, int(paragraph_data.get("indentation_level", 1)) - 1)
                    p.level = level
                    logger.debug(f"Bullet[{idx}]: level={level} text={p.text!r}")

        except Exception as e:
            logger.error(f"Failed to create content slide: {e}")
            raise

    # def save(self) -> io.BytesIO:
    #     """Save presentation to BytesIO object."""
    #     try:
    #         logger.info("Saving PowerPoint to memory buffer")
    #         file_like_object = io.BytesIO()
    #         self.presentation.save(file_like_object)
    #         file_like_object.seek(0)
    #         return file_like_object
    #     except Exception as e:
    #         logger.error(f"Failed to save presentation: {e}")
    #         raise
    def save(self) -> str:
        """Save presentation to disk and return file path."""
        try:
            import os
            import uuid

            OUTPUT_DIR = "/app/output"
            os.makedirs(OUTPUT_DIR, exist_ok=True)

            # Generate filename
            filename = f"presentation_{uuid.uuid4().hex}.pptx"
            file_path = os.path.join(OUTPUT_DIR, filename)

            # Save to disk
            self.presentation.save(file_path)

            logger.info(f"Saved PowerPoint to {file_path}")
            return file_path

        except Exception as e:
            logger.error(f"Failed to save presentation: {e}")
            raise

